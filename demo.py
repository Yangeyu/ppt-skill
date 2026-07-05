"""Build a real deck end-to-end and self-verify it is native + editable."""
from __future__ import annotations
import sys
from pathlib import Path
import zipfile
from pptx import Presentation
from ppt_engine.build import Engine
from ppt_engine.selfcheck import check_layout, render_preview
from ppt_engine.ir import (
    Deck, DeckMeta, Slide,
    CoverData, SectionData, KpiData, Stat, BulletsData, Bullet,
    ChartData, Series, TocData, TocItem, TwoColData, CompareData, Column,
    ProcessData, Step, QuoteData, ClosingData,
)

DECK = Deck(
    meta=DeckMeta(title="2026 上半年增长复盘"),
    theme="editorial",
    slides=[
        Slide(data=CoverData(
            eyebrow="Growth Review",
            title="2026 上半年增长复盘",
            subtitle="增长团队 · Q2 季度业务汇报与下半年规划",
            footer="Binyang  ·  2026-06-26")),
        Slide(data=TocData(
            eyebrow="Agenda", title="目录",
            items=[TocItem(title="业绩概览"), TocItem(title="增长抓手"),
                   TocItem(title="增长方法论"), TocItem(title="趋势洞察"),
                   TocItem(title="成效与挑战"), TocItem(title="下半年规划")])),
        Slide(data=KpiData(
            eyebrow="Overview", title="核心指标概览",
            stats=[
                Stat(value="+42%", label="GMV 同比", delta="较 Q1 +9pct", delta_dir="up"),
                Stat(value="1.2M", label="月活用户", delta="环比 +14%", delta_dir="up"),
                Stat(value="3.8%", label="支付转化率", delta="较 Q1 +0.6pct", delta_dir="up"),
                Stat(value="-18%", label="获客成本", delta="渠道优化", delta_dir="flat"),
            ])),
        Slide(data=BulletsData(
            eyebrow="What worked", title="本季度增长抓手",
            bullets=[
                Bullet(text="渠道组合优化:压缩低效投放、放大高 ROI 渠道,CAC 下降 18%", emphasis=True),
                Bullet(text="新用户首日留存从 41% 提升至 55%,靠的是新手引导与首单激励重构"),
                Bullet(text="会员体系上线,付费会员贡献了 GMV 增量的 27%"),
                Bullet(text="搜索与推荐打通,商详页转化率提升 0.6 个百分点"),
                Bullet(text="供给侧引入 120 个新品牌,丰富了中高客单价心智"),
            ])),
        Slide(data=ProcessData(
            eyebrow="Method", title="增长方法论:四步闭环",
            steps=[
                Step(title="洞察", desc="数据埋点 + 用户访谈定位机会", icon="target"),
                Step(title="实验", desc="A/B 快速验证假设", icon="zap"),
                Step(title="放大", desc="跑通的策略加大投入", icon="trending-up"),
                Step(title="沉淀", desc="固化为可复用的 SOP", icon="layers"),
            ])),
        Slide(data=ChartData(
            eyebrow="Trend", title="GMV 月度趋势(百万元)",
            chart_type="column",
            categories=["1月", "2月", "3月", "4月", "5月", "6月"],
            series=[Series(name="GMV", values=[82, 78, 95, 110, 128, 146])],
            takeaway="6 月 GMV 创历史新高,环比 +14%,Q2 增速明显快于 Q1。")),
        Slide(data=CompareData(
            eyebrow="Review", title="成效与挑战",
            left=Column(heading="做对了", points=[
                "高 ROI 渠道结构跑通", "会员体系验证了付费意愿", "留存曲线整体上移"]),
            right=Column(heading="待改进", points=[
                "腰部用户活跃度仍偏低", "供给丰富度不足以支撑客单价", "履约时效投诉环比上升"]))),
        Slide(data=SectionData(
            number="06", title="下半年规划", subtitle="从增长提速转向高质量增长")),
        Slide(data=TwoColData(
            eyebrow="Plan", title="两个主攻方向",
            left=Column(heading="盈利质量", points=[
                "把单位经济模型做正,目标毛利率 +3pct", "优化履约与退货成本结构", "高价值供给的专项扶持"]),
            right=Column(heading="用户分层", points=[
                "沉淀 RFM 模型与分层策略", "对高价值用户做专属权益", "腰部用户的激活与召回"]))),
        Slide(data=QuoteData(
            quote="增长不是把漏斗灌满,而是让每一层都更值得留下。",
            attribution="—— 增长团队 2026 年中复盘")),
        Slide(data=ClosingData(
            title="谢谢",
            subtitle="期待下半年与各业务线一起把高质量增长做实",
            contact="增长团队  ·  growth@company.com")),
    ],
)


def verify(path: str):
    prs = Presentation(path)
    print(f"\n=== 结构自检: {path} ===")
    print(f"画布 {prs.slide_width / 914400:.2f}×{prs.slide_height / 914400:.2f} in · 共 {len(prs.slides)} 页")
    for i, slide in enumerate(prs.slides, 1):
        n_text = n_shape = n_chart = n_pic = 0
        for sh in slide.shapes:
            if sh.has_chart:
                n_chart += 1
            elif sh.shape_type == 13:  # PICTURE
                n_pic += 1
            elif sh.has_text_frame and sh.text_frame.text.strip():
                n_text += 1
            else:
                n_shape += 1
        extra = f" · 图片 {n_pic}" if n_pic else ""
        extra += f" · 原生图表 {n_chart}" if n_chart else ""
        print(f"  slide {i:2d}: 文本 {n_text:2d} · 形状 {n_shape:2d}{extra}")
    # embedded fonts?
    z = zipfile.ZipFile(path)
    fonts = [n for n in z.namelist() if n.startswith("ppt/fonts/")]
    embedded = "embedTrueTypeFonts" in z.read("ppt/presentation.xml").decode()
    print(f"内嵌字体: {len(fonts)} 个 fntdata 部件, embedTrueTypeFonts={'on' if embedded else 'off'} "
          f"(总 {sum(len(z.read(f)) for f in fonts) // 1024} KB)")


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "growth_review.pptx"
    shots = sys.argv[2] if len(sys.argv) > 2 else None
    prims = Engine().build(DECK, out, screenshot_dir=shots)
    total = sum(len(p) for p in prims)
    print(f"已生成 {out} — {len(prims)} 页, {total} 个定位原语")
    # self-check layer 1: structural invariants (no render)
    issues = check_layout(prims)
    print(f"结构自检(几何层): {'零问题 ✓' if not issues else issues}")
    verify(out)

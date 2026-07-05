"""Renderer — translate resolved primitives into NATIVE, editable pptx objects.
Text -> textbox, rect -> autoshape, chart -> native chart, icon -> picture.
Token colors are emitted as theme refs (schemeClr) so the deck stays recolorable."""
from __future__ import annotations
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml
from .units import emu, parse_color, SLIDE_W_EMU, SLIDE_H_EMU
from .oox_theme import inject_theme

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT,
         "start": PP_ALIGN.LEFT, "end": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
CHART_T = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
           "line": XL_CHART_TYPE.LINE_MARKERS}
ZRANK = {"rect": 0, "line": 0, "image": 5, "icon": 6, "chart": 5, "text": 10}
# families that carry CJK glyphs — used to split a:latin (Latin face) from
# a:ea (East-Asian face) so a heavy Latin display + a CJK black can coexist in one run
CJK_FAMILIES = {"Noto Sans SC", "Noto Serif SC", "Noto Sans SC Black", "PingFang SC",
                "Songti SC", "STSong", "Hiragino Sans GB", "Microsoft YaHei", "SimHei"}
TEXT_SLOT = {
    "tx1": MSO_THEME_COLOR.TEXT_1, "bg1": MSO_THEME_COLOR.BACKGROUND_1,
    "tx2": MSO_THEME_COLOR.TEXT_2, "bg2": MSO_THEME_COLOR.BACKGROUND_2,
    "accent1": MSO_THEME_COLOR.ACCENT_1, "accent2": MSO_THEME_COLOR.ACCENT_2,
}


def _zkey(item):
    i, p = item
    z = p.get("z")
    return (z if z is not None else ZRANK.get(p["kind"], 0), i)


# ---- OOXML fill helpers --------------------------------------------------
def _clr_inner(theme, hexv, alpha=1.0):
    a = f'<a:alpha val="{int(alpha * 100000)}"/>' if alpha < 1 else ''
    slot = theme.color_slot().get(hexv)
    if slot:
        return f'<a:schemeClr val="{slot}">{a}</a:schemeClr>'
    return f'<a:srgbClr val="{hexv}">{a}</a:srgbClr>'


def _remove_fills(spPr):
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill"):
        e = spPr.find(qn(tag))
        if e is not None:
            spPr.remove(e)


def _insert_fill(spPr, fill_el):
    ln = spPr.find(qn("a:ln"))
    (ln.addprevious if ln is not None else spPr.append)(fill_el)


def _set_solid(shape, theme, hexv, alpha=1.0):
    spPr = shape._element.spPr
    _remove_fills(spPr)
    _insert_fill(spPr, parse_xml(f'<a:solidFill {nsdecls("a")}>{_clr_inner(theme, hexv, alpha)}</a:solidFill>'))


def _set_nofill(shape):
    spPr = shape._element.spPr
    _remove_fills(spPr)
    _insert_fill(spPr, parse_xml(f'<a:noFill {nsdecls("a")}/>'))


def _set_gradient(shape, c1, c2, angle):
    spPr = shape._element.spPr
    _remove_fills(spPr)
    _insert_fill(spPr, parse_xml(
        f'<a:gradFill {nsdecls("a")} rotWithShape="1"><a:gsLst>'
        f'<a:gs pos="0"><a:srgbClr val="{c1}"/></a:gs>'
        f'<a:gs pos="100000"><a:srgbClr val="{c2}"/></a:gs>'
        f'</a:gsLst><a:lin ang="{int(angle * 60000)}" scaled="1"/></a:gradFill>'))


def _set_run_fonts(run, latin, ea):
    run.font.name = latin
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            rPr.append(parse_xml(f'<{tag} {nsdecls("a")} typeface="{ea}"/>'))
        else:
            el.set("typeface", ea)


# ---- primitives ----------------------------------------------------------
def _add_rect(slide, p, theme):
    x, y, w, h = emu(p["x"]), emu(p["y"]), emu(p["w"]), emu(p["h"])
    radius = p.get("radiusPx", 0) or 0
    short = min(p["w"], p["h"])
    is_circle = radius >= short / 2 - 0.5 and abs(p["w"] - p["h"]) <= 2
    if is_circle:
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    elif radius:
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        try:
            shp.adjustments[0] = max(0.0, min(0.5, radius / short))
        except Exception:
            pass
    else:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)

    grad = p.get("grad")
    if grad:
        c1, c2, ang = grad.split(",")
        _set_gradient(shp, c1, c2, float(ang))
    else:
        hexv, alpha = parse_color(p.get("fill"))
        _set_nofill(shp) if (hexv is None or alpha == 0) else _set_solid(shp, theme, hexv, alpha)

    bw = p.get("borderWidthPx", 0) or 0
    bc, _ = parse_color(p.get("borderColor")) if p.get("borderColor") else (None, 0)
    if bw and bc:
        shp.line.color.rgb = RGBColor.from_string(bc)
        shp.line.width = Pt(bw * 0.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False


def _add_text(slide, p, theme):
    # add a little width slack at the wrap edge so a sub-pixel metric diff
    # between browser and PowerPoint can't flip the line-wrap (keep start aligned)
    align = p.get("align", "left")
    pad = emu(8)
    x, y, w, h = emu(p["x"]), emu(p["y"]), emu(p["w"]), emu(p["h"]) + emu(2)
    if align == "center":
        x -= pad // 2
        w += pad
    elif align in ("right", "end"):
        x -= pad
        w += pad
    else:
        w += pad
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = not p.get("nowrap", False)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    try:
        tf.auto_size = None
    except Exception:
        pass
    para = tf.paragraphs[0]
    para.alignment = ALIGN.get(p.get("align", "left"), PP_ALIGN.LEFT)
    fs = p["fontSizePx"]
    lh = p.get("lineHeightPx", fs * 1.2)
    if p.get("nowrap"):
        # exact spacing (pt) = the browser-measured line box. Multiple spacing
        # scales the FONT's own line height (CJK ~1.4em), which draws display
        # glyphs tens of px lower than measured and overlaps the block below.
        para.line_spacing = Pt(lh * 0.75)
    else:
        para.line_spacing = max(0.8, lh / fs)

    run = para.add_run()
    run.text = p["text"].upper() if p.get("upper") else p["text"]
    f = run.font
    f.size = Pt(fs * 0.75)
    f.bold = p.get("fontWeight", 400) >= 600
    f.italic = bool(p.get("italic"))
    hexv, _ = parse_color(p.get("color"))
    slot = theme.color_slot().get(hexv) if hexv else None
    if slot in TEXT_SLOT:
        f.color.theme_color = TEXT_SLOT[slot]
    elif hexv:
        f.color.rgb = RGBColor.from_string(hexv)
    fams = p.get("families") or ([p["family"]] if p.get("family") else [theme.body_font])
    latin = fams[0]
    ea = next((f for f in fams if f in CJK_FAMILIES), latin)
    _set_run_fonts(run, latin, ea)
    ls = p.get("letterSpacingPx", 0) or 0
    if ls:
        run._r.get_or_add_rPr().set("spc", str(int(ls * 0.75 * 100)))


def _add_picture(slide, p):
    if p.get("src"):
        slide.shapes.add_picture(p["src"], emu(p["x"]), emu(p["y"]), emu(p["w"]), emu(p["h"]))


C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def _chart_transparent(chart):
    """Make chart area + plot area transparent so the slide background shows
    through (no white box on a paper deck). Best-effort."""
    try:
        cs = chart._chartSpace
        nofill = lambda: parse_xml(
            f'<c:spPr xmlns:c="{C_NS}" {nsdecls("a")}><a:noFill/>'
            f'<a:ln><a:noFill/></a:ln></c:spPr>')
        chart_el = cs.find(qn("c:chart"))
        chart_el.addnext(nofill())                      # chartSpace fill
        plot = chart_el.find(qn("c:plotArea"))
        plot.append(nofill())                           # plotArea fill
    except Exception:
        pass


def _add_chart(slide, p, theme):
    c = p["chart"]
    ctype = c.get("type", "column")
    data = CategoryChartData()
    data.categories = c["categories"]
    for s in c["series"]:
        data.add_series(s["name"], s["values"])
    gf = slide.shapes.add_chart(CHART_T.get(ctype, XL_CHART_TYPE.COLUMN_CLUSTERED),
                                emu(p["x"]), emu(p["y"]), emu(p["w"]), emu(p["h"]), data)
    chart = gf.chart
    chart.has_title = False
    multi = len(c["series"]) > 1
    chart.has_legend = multi
    if multi:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    chart.font.size = Pt(12)
    _chart_transparent(chart)
    for i, series in enumerate(chart.plots[0].series):
        hexv = theme.chart_palette[i % len(theme.chart_palette)]
        if ctype == "line":
            series.format.line.color.rgb = RGBColor.from_string(hexv)
            series.format.line.width = Pt(2.5)
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor.from_string(hexv)
            series.format.line.fill.background()


def render_deck(deck, slides_prims, theme, out_path):
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(SLIDE_H_EMU)
    inject_theme(prs, theme)
    blank = prs.slide_layouts[6]
    for prims in slides_prims:
        slide = prs.slides.add_slide(blank)
        for _, p in sorted(enumerate(prims), key=_zkey):
            k = p["kind"]
            if k == "text":
                _add_text(slide, p, theme)
            elif k == "rect":
                _add_rect(slide, p, theme)
            elif k in ("image", "icon"):
                _add_picture(slide, p)
            elif k == "chart":
                _add_chart(slide, p, theme)
    prs.save(out_path)
